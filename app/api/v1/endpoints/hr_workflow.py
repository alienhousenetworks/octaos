"""API Endpoints for OctaOS Dynamic HR Document Intelligence & Workflow Builder."""
import hashlib
import io
import uuid
from datetime import datetime, timezone
from typing import List, Dict, Any, Optional
from pydantic import BaseModel, Field

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, status
from sqlalchemy.orm import Session

from app.api import deps
from app.models.base import User
from app.models.verticals import Candidate
from app.models.hr import (
    TenantDocumentTemplate,
    TemplateVersion,
    TemplateVariableSchema,
    HRWorkflowDefinition,
    HRWorkflowVersion,
    HRWorkflowImportRecord,
    HRWorkflowInstance,
    HRWorkflowTask,
    CandidateFieldValue,
    CandidateEvidence,
    DocumentDraft,
    DocumentReviewComment,
    ScheduledDelivery,
    Application,
    JobRequisition,
)
from app.services.hr.workflow.import_service import WorkflowImportService
from app.schemas.hr_workflow import (
    TemplateAnalysisResponse,
    TemplatePublishRequest,
    WorkflowDefinitionCreate,
    WorkflowDefinitionUpdate,
    NodeTypeMetadata,
    WorkflowTestRunRequest,
    WorkflowTestRunResponse,
    WorkflowGraphSchema,
    AIPipelineGenerateRequest,
    AIPipelineGenerateResponse,
    DraftPreviewResponse,
    DraftEditRequest,
    DraftApprovalRequest,
    DraftScheduleRequest,
    ReconciliationSummary,
    AddReviewCommentRequest,
    OnboardingInstanceCreateRequest,
    OnboardingInstanceUpdateRequest,
    OnboardingInstanceSummary,
    RequiredDocumentItem,
)
from app.services.hr.template_hub.parser import DocumentTemplateParser
from app.services.hr.template_hub.intelligence import TemplateIntelligenceAnalyzer
from app.services.hr.workflow.engine import DynamicWorkflowEngine
from app.services.hr.candidate_portal.service import CandidatePortalService
from app.services.hr.extraction.reconciliation import EvidenceReconciliationService
from app.services.hr.review_studio.renderer import ReviewStudioRenderer
from app.services.hr.delivery.scheduler import DeliverySchedulerService
from app.services.hr.document_store import DocumentStore
from app.services.hr.document_verification_service import DocumentVerificationService
from app.services.hr.employee_provisioning_service import EmployeeProvisioningService


router = APIRouter()


# --- TEMPLATE HUB ENDPOINTS ---

@router.post("/templates/upload", response_model=TemplateAnalysisResponse)
async def upload_template(
    file: UploadFile = File(...),
    name: str = Form(...),
    category: str = Form("onboarding"),
    folder_path: str = Form("/"),
    tags: str = Form(""),
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Uploads a DOCX or HTML template, calculates SHA-256, parses structure, and proposes canonical mappings."""
    tenant_id = current_user.tenant_id
    content_bytes = await file.read()
    filename = file.filename or "template.docx"
    doc_format = "docx" if filename.lower().endswith(".docx") else "html"
    content_sha = hashlib.sha256(content_bytes).hexdigest()

    # Parse structure and placeholders
    paragraphs_sample: List[str] = []
    if doc_format == "docx":
        placeholders, meta = DocumentTemplateParser.parse_docx(content_bytes)
        try:
            import docx
            doc = docx.Document(io.BytesIO(content_bytes))
            paragraphs_sample = [p.text for p in doc.paragraphs if p.text]
        except Exception:
            paragraphs_sample = []
    else:
        html_str = content_bytes.decode("utf-8", errors="ignore")
        placeholders, meta = DocumentTemplateParser.parse_html(html_str)
        paragraphs_sample = [s.strip() for s in html_str.split("</p>") if s.strip()]

    # Run AI intelligence analyzer
    proposed_mappings, detected_clauses = TemplateIntelligenceAnalyzer.analyze_template(
        placeholders=placeholders,
        paragraphs_text=paragraphs_sample,
    )
    deduced_requirements = TemplateIntelligenceAnalyzer.deduce_required_candidate_documents(
        placeholders=placeholders,
        paragraphs_text=paragraphs_sample,
        detected_clauses=detected_clauses,
    )
    meta["deduced_requirements"] = deduced_requirements

    # Store file binary safely
    storage_key = DocumentStore.save(
        tenant_id=tenant_id,
        filename=filename,
        content=content_bytes,
        prefix="templates",
    )

    # Persist or update template
    parsed_tags = [t.strip() for t in tags.split(",") if t.strip()]
    template = db.query(TenantDocumentTemplate).filter_by(
        tenant_id=tenant_id,
        name=name,
    ).first()

    if not template:
        template = TenantDocumentTemplate(
            tenant_id=tenant_id,
            document_type=category,
            category=category,
            name=name,
            folder_path=folder_path,
            tags=parsed_tags,
        )
        db.add(template)
        db.flush()

    # Create new version
    version_count = db.query(TemplateVersion).filter_by(template_id=template.id).count()
    template_version = TemplateVersion(
        template_id=template.id,
        tenant_id=tenant_id,
        version=version_count + 1,
        format=doc_format,
        status="analyzed",
        storage_key=storage_key,
        content_sha256=content_sha,
        raw_content=content_bytes.decode("utf-8", errors="ignore") if doc_format == "html" else None,
        docx_structure=meta,
        detected_clauses=[c.model_dump() for c in detected_clauses],
        proposed_mappings=[m.model_dump() for m in proposed_mappings],
    )
    db.add(template_version)
    db.commit()
    db.refresh(template_version)

    return TemplateAnalysisResponse(
        template_id=template.id,
        version_id=template_version.id,
        format=doc_format,
        content_sha256=content_sha,
        placeholders=placeholders,
        proposed_mappings=proposed_mappings,
        detected_clauses=detected_clauses,
        deduced_requirements=deduced_requirements,
        table_count=meta.get("table_count", 0),
        paragraph_count=meta.get("paragraph_count", 0),
    )


@router.post("/templates/{version_id}/publish")
def publish_template(
    version_id: str,
    req: TemplatePublishRequest,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Approves canonical mappings and publishes the template version for workflows."""
    version = db.query(TemplateVersion).filter_by(
        id=version_id,
        tenant_id=current_user.tenant_id,
    ).first()
    if not version:
        raise HTTPException(status_code=404, detail="Template version not found")

    # Record approved variable schemas
    for m in req.approved_mappings:
        var_schema = TemplateVariableSchema(
            template_version_id=version.id,
            variable_name=m.canonical_key,
            data_type=m.data_type,
            is_required=m.is_required,
            source_field=m.canonical_key,
            sensitivity="confidential",
            editable_by_recruiter=True,
        )
        db.add(var_schema)

    version.status = "published"
    version.published_at = datetime.now(timezone.utc)
    version.published_by_user_id = current_user.id
    db.commit()
    return {"status": "published", "version_id": version.id}


@router.get("/templates")
def list_templates(
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Lists all document templates and their latest versions for the tenant."""
    templates = db.query(TenantDocumentTemplate).filter_by(
        tenant_id=current_user.tenant_id,
    ).all()

    result = []
    for t in templates:
        latest_ver = (
            db.query(TemplateVersion)
            .filter_by(template_id=t.id)
            .order_by(TemplateVersion.version.desc())
            .first()
        )
        result.append({
            "id": t.id,
            "name": t.name,
            "category": t.category,
            "folder_path": t.folder_path,
            "tags": t.tags or [],
            "latest_version": latest_ver.version if latest_ver else 1,
            "status": latest_ver.status if latest_ver else "draft",
            "format": latest_ver.format if latest_ver else "docx",
            "created_at": t.created_at.isoformat() if t.created_at else None,
        })
    return result


# --- WORKFLOW BUILDER ENDPOINTS ---

@router.post("/definitions")
def create_workflow_definition(
    req: WorkflowDefinitionCreate,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Creates a new dynamic workflow graph definition."""
    is_valid, errors = DynamicWorkflowEngine.validate_graph(req.graph_schema)
    if not is_valid:
        raise HTTPException(status_code=400, detail={"errors": errors})

    wf = HRWorkflowDefinition(
        tenant_id=current_user.tenant_id,
        name=req.name,
        description=req.description,
        graph_schema=req.graph_schema.model_dump(),
    )
    db.add(wf)
    db.commit()
    db.refresh(wf)
    return {"id": wf.id, "name": wf.name, "version": wf.version}


class WorkflowImportRequest(BaseModel):
    target_folder_path: str = "/"
    target_name: Optional[str] = None
    target_department_id: Optional[str] = None
    clone_templates: bool = True


@router.get("/folders", response_model=List[str])
def list_workflow_folders(
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Returns all available company folders and departments holding workflows or templates."""
    return WorkflowImportService.get_all_folders(db, current_user.tenant_id)


@router.post("/definitions/{definition_id}/import-to-folder")
def import_workflow_to_folder(
    definition_id: str,
    req: WorkflowImportRequest,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """
    Imports and clones an existing workflow and its document templates across folders/departments.
    Remaps internal node IDs and edges, clones templates into target folder, and preserves provenance.
    Saves as an editable draft.
    """
    try:
        new_wf = WorkflowImportService.clone_workflow(
            db=db,
            tenant_id=current_user.tenant_id,
            user_id=current_user.id,
            source_definition_id=definition_id,
            target_folder_path=req.target_folder_path,
            target_name=req.target_name,
            target_department_id=req.target_department_id,
            clone_templates=req.clone_templates,
        )
        return {
            "status": "imported",
            "id": new_wf.id,
            "name": new_wf.name,
            "folder_path": new_wf.folder_path,
            "version": new_wf.version,
            "is_active": new_wf.is_active,
        }
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Import failed: {str(e)}")


@router.get("/definitions")
def list_workflow_definitions(
    folder_path: Optional[str] = None,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Lists all workflow definitions for the current tenant, with optional folder filter."""
    query = db.query(HRWorkflowDefinition).filter_by(tenant_id=current_user.tenant_id)
    if folder_path:
        query = query.filter(HRWorkflowDefinition.folder_path == folder_path)
    workflows = query.all()
    if not workflows:
        default_graph = {
            "nodes": [
                {
                    "id": "step_intake",
                    "name": "1. Candidate Information Intake",
                    "type": "start",
                    "position": {"x": 60, "y": 180},
                    "config": {
                        "trigger_type": "manual",
                        "auto_start_on_candidate_creation": True,
                        "intake_fields": [
                            {"key": "candidate.full_name", "label": "Full Legal Name", "is_required": True},
                            {"key": "candidate.email", "label": "Email Address", "is_required": True},
                            {"key": "candidate.phone", "label": "Phone Number", "is_required": True},
                            {"key": "candidate.current_address", "label": "Residential Address", "is_required": True},
                            {"key": "candidate.pan_number", "label": "PAN / Tax ID", "is_required": True},
                        ],
                    },
                },
                {
                    "id": "step_docs",
                    "name": "2. Required Document Evidence",
                    "type": "candidate_upload",
                    "position": {"x": 320, "y": 180},
                    "config": {
                        "required_documents": [
                            {"title": "Government Photo ID", "category": "passport", "is_mandatory": True, "description": "Passport, Aadhaar, or National ID"},
                            {"title": "Tax Proof / PAN Card", "category": "tax_form", "is_mandatory": True, "description": "PAN card or national tax proof"},
                            {"title": "Educational Certificate", "category": "degree_certificate", "is_mandatory": True, "description": "Highest degree certificate or transcript"},
                        ],
                        "allowed_formats": ["application/pdf", "image/jpeg", "image/png"],
                        "deadline_days": 5,
                    },
                },
                {
                    "id": "step_branch",
                    "name": "3. Employment Type Routing",
                    "type": "condition_branch",
                    "position": {"x": 580, "y": 180},
                    "config": {
                        "branch_attribute": "employment_type",
                        "branches": [
                            {"label": "Permanent FTE", "condition": "context.employment_type == 'permanent'"},
                            {"label": "Independent Contractor", "condition": "context.employment_type == 'contractor'"},
                        ],
                    },
                },
                {
                    "id": "step_offer",
                    "name": "4. Offer & Agreement Generation",
                    "type": "generate_document",
                    "position": {"x": 840, "y": 180},
                    "config": {
                        "template_name": "Executive Employment Agreement",
                        "output_format": "both",
                        "missing_data_policy": "route_to_hr",
                        "field_mappings": {
                            "candidate_name": "candidate.full_name",
                            "annual_ctc": "company_fields.annual_ctc",
                            "role_title": "company_fields.role_title",
                        },
                    },
                },
                {
                    "id": "step_approval",
                    "name": "5. HR Sign-off & Cryptographic Binding",
                    "type": "approval",
                    "position": {"x": 1100, "y": 180},
                    "config": {
                        "required_role": "hr_director",
                        "approval_policy": "strict_sha_binding",
                        "timeout_hours": 48,
                    },
                },
                {
                    "id": "step_complete",
                    "name": "6. Dispatch & Complete Hire",
                    "type": "end",
                    "position": {"x": 1360, "y": 180},
                    "config": {
                        "completion_state": "hired",
                        "sync_to_payroll": True,
                        "notify_participants": True,
                    },
                },
            ],
            "edges": [
                {"id": "e_intake_docs", "source": "step_intake", "target": "step_docs"},
                {"id": "e_docs_branch", "source": "step_docs", "target": "step_branch"},
                {"id": "e_branch_offer", "source": "step_branch", "target": "step_offer"},
                {"id": "e_offer_approval", "source": "step_offer", "target": "step_approval"},
                {"id": "e_approval_complete", "source": "step_approval", "target": "step_complete"},
            ],
        }
        seeded_wf = HRWorkflowDefinition(
            tenant_id=current_user.tenant_id,
            name="Standard Autonomous Onboarding Flow",
            description="Autonomous onboarding pipeline with smart candidate intake, verification evidence, routing, legal generation, and HR sign-off.",
            is_active=True,
            graph_schema=default_graph,
        )
        db.add(seeded_wf)
        db.commit()
        db.refresh(seeded_wf)
        workflows = [seeded_wf]

    return [
        {
            "id": w.id,
            "name": w.name,
            "description": w.description,
            "folder_path": w.folder_path or "/",
            "department_id": w.department_id,
            "published_version_id": w.published_version_id,
            "is_active": w.is_active,
            "version": w.version,
            "node_count": len((w.graph_schema or {}).get("nodes", [])),
            "created_at": w.created_at.isoformat() if w.created_at else None,
        }
        for w in workflows
    ]


@router.get("/node-registry", response_model=List[NodeTypeMetadata])
def get_node_registry():
    """Returns the comprehensive registry of all 15 configurable node types."""
    return DynamicWorkflowEngine.get_node_registry()


@router.post("/ai-generate", response_model=AIPipelineGenerateResponse)
def ai_generate_workflow(
    req: AIPipelineGenerateRequest,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Synthesizes a proposed DAG workflow from a natural-language prompt.
    Validates the generated graph against the shared WorkflowGraphContract schema.
    Does NOT publish or execute candidate actions automatically.
    """
    return DynamicWorkflowEngine.generate_workflow_graph_from_prompt(
        prompt=req.prompt,
        job_role=req.job_role,
        tenant_id=current_user.tenant_id,
        db=db,
    )


@router.get("/definitions/{definition_id}")
def get_workflow_definition(
    definition_id: str,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Retrieves full workflow definition with its graph schema for the visual canvas."""
    wf = db.query(HRWorkflowDefinition).filter_by(
        id=definition_id,
        tenant_id=current_user.tenant_id,
    ).first()
    if not wf:
        raise HTTPException(status_code=404, detail="Workflow definition not found")

    return {
        "id": wf.id,
        "name": wf.name,
        "description": wf.description,
        "is_active": wf.is_active,
        "version": wf.version,
        "graph_schema": wf.graph_schema or {"nodes": [], "edges": []},
        "created_at": wf.created_at.isoformat() if wf.created_at else None,
        "updated_at": wf.updated_at.isoformat() if wf.updated_at else None,
    }


@router.put("/definitions/{definition_id}")
def update_workflow_definition(
    definition_id: str,
    req: WorkflowDefinitionUpdate,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Updates a workflow definition and its visual graph schema."""
    wf = db.query(HRWorkflowDefinition).filter_by(
        id=definition_id,
        tenant_id=current_user.tenant_id,
    ).first()
    if not wf:
        raise HTTPException(status_code=404, detail="Workflow definition not found")

    if req.graph_schema:
        is_valid, errors = DynamicWorkflowEngine.validate_graph(req.graph_schema)
        if not is_valid:
            raise HTTPException(status_code=400, detail={"errors": errors})
        wf.graph_schema = req.graph_schema.model_dump()

    if req.name is not None:
        wf.name = req.name
    if req.description is not None:
        wf.description = req.description
    if req.is_active is not None:
        wf.is_active = req.is_active

    db.commit()
    db.refresh(wf)
    return {
        "id": wf.id,
        "name": wf.name,
        "version": wf.version,
        "is_active": wf.is_active,
        "updated_at": wf.updated_at.isoformat() if wf.updated_at else None,
    }


@router.post("/definitions/{definition_id}/publish")
def publish_workflow_definition(
    definition_id: str,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Validates and publishes an immutable version of the workflow definition."""
    wf = db.query(HRWorkflowDefinition).filter_by(
        id=definition_id,
        tenant_id=current_user.tenant_id,
    ).first()
    if not wf:
        raise HTTPException(status_code=404, detail="Workflow definition not found")

    # 1. Validate graph structure before publishing
    schema = WorkflowGraphSchema(**(wf.graph_schema or {"nodes": [], "edges": []}))
    is_valid, errors = DynamicWorkflowEngine.validate_graph(schema)
    if not is_valid:
        raise HTTPException(status_code=400, detail={"detail": "Cannot publish invalid workflow", "errors": errors})

    new_version_num = (wf.version or 1) + 1

    # 2. Persist immutable published version
    ver_record = HRWorkflowVersion(
        definition_id=wf.id,
        tenant_id=current_user.tenant_id,
        version=new_version_num,
        graph_schema=wf.graph_schema or {},
        published_by_user_id=current_user.id,
        change_summary="Published from visual workflow studio",
    )
    db.add(ver_record)
    db.flush()

    wf.version = new_version_num
    wf.published_version_id = ver_record.id
    wf.is_active = True
    db.commit()
    return {
        "status": "published",
        "id": wf.id,
        "version": wf.version,
        "published_version_id": ver_record.id,
    }


@router.delete("/definitions/{definition_id}")
def delete_workflow_definition(
    definition_id: str,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Deletes or deactivates a workflow definition."""
    wf = db.query(HRWorkflowDefinition).filter_by(
        id=definition_id,
        tenant_id=current_user.tenant_id,
    ).first()
    if not wf:
        raise HTTPException(status_code=404, detail="Workflow definition not found")

    wf.is_active = False
    db.commit()
    return {"status": "deactivated", "id": wf.id}


@router.post("/definitions/test-run", response_model=WorkflowTestRunResponse)
def test_run_workflow(req: WorkflowTestRunRequest):
    """Simulates a test run of a workflow DAG against sample candidate context."""
    return DynamicWorkflowEngine.evaluate_traversal(
        graph=req.graph_schema,
        context=req.sample_context,
    )


@router.post("/definitions/{definition_id}/instantiate")
def instantiate_workflow(
    definition_id: str,
    candidate_id: str = Form(...),
    application_id: str = Form(...),
    context_data: Optional[Dict[str, Any]] = None,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Launches a dynamic workflow instance for a candidate application and generates candidate portal link."""
    wf_def = db.query(HRWorkflowDefinition).filter_by(
        id=definition_id,
        tenant_id=current_user.tenant_id,
    ).first()
    if not wf_def:
        raise HTTPException(status_code=404, detail="Workflow definition not found")

    raw_token, token_hash, expires_at = CandidatePortalService.generate_token_pair(expires_in_days=7)

    ctx = dict(context_data or {})
    ctx["raw_portal_token"] = raw_token

    instance = HRWorkflowInstance(
        tenant_id=current_user.tenant_id,
        workflow_definition_id=wf_def.id,
        application_id=application_id,
        candidate_id=candidate_id,
        current_state="intake_pending",
        candidate_token_hash=token_hash,
        token_expires_at=expires_at,
        context_data=ctx,
    )
    db.add(instance)
    db.flush()

    # Create initial tasks from graph
    nodes = (wf_def.graph_schema or {}).get("nodes", [])
    for node in nodes:
        task = HRWorkflowTask(
            tenant_id=current_user.tenant_id,
            instance_id=instance.id,
            node_id=node.get("id"),
            task_type=node.get("type", "candidate_intake"),
            status="pending",
        )
        db.add(task)

    db.commit()

    return {
        "instance_id": instance.id,
        "candidate_portal_token": raw_token,
        "candidate_portal_url": f"/candidate-portal/{raw_token}",
        "expires_at": expires_at.isoformat(),
        "status": instance.current_state,
    }


def _serialize_instance_summary(db: Session, instance: HRWorkflowInstance) -> OnboardingInstanceSummary:
    candidate = db.query(Candidate).filter_by(id=instance.candidate_id).first()
    app = db.query(Application).filter_by(id=instance.application_id).first() if instance.application_id else None
    req = db.query(JobRequisition).filter_by(id=app.requisition_id).first() if (app and app.requisition_id) else None
    role_title = (
        (instance.context_data or {}).get("company_fields", {}).get("role_title")
        or (instance.context_data or {}).get("role_title")
        or (req.title if req else None)
        or (candidate.role if candidate else None)
        or "Employee"
    )

    ctx = instance.context_data or {}
    req_docs = CandidatePortalService.get_required_documents(ctx)
    evidence_records = db.query(CandidateEvidence).filter_by(instance_id=instance.id).all()
    uploaded_evidence = [
        {
            "id": e.id,
            "document_name": e.document_name,
            "document_category": e.document_category,
            "file_sha256": e.file_sha256,
            "uploaded_at": e.uploaded_at.isoformat() if e.uploaded_at else "",
        }
        for e in evidence_records
    ]
    field_records = db.query(CandidateFieldValue).filter_by(instance_id=instance.id).all()
    field_values = {f.canonical_key: f.resolved_value or f.candidate_provided_value for f in field_records}

    synthesized_fields = CandidatePortalService.synthesize_intake_questionnaire([])
    missing_company, missing_candidate = CandidatePortalService.audit_missing_items(
        context=ctx,
        synthesized_fields=synthesized_fields,
        required_documents=req_docs,
        uploaded_evidence=uploaded_evidence,
        field_values=field_values,
    )

    raw_token = ctx.get("raw_portal_token", "")
    portal_url = f"/candidate-portal/{raw_token}" if raw_token else f"/candidate-portal/{instance.id}"

    total_checks = max(1, len(req_docs) + 4)
    completed_checks = (len(req_docs) - len(missing_candidate)) + (4 - len(missing_company))
    progress = int(max(0, min(100, (completed_checks / total_checks) * 100)))

    return OnboardingInstanceSummary(
        id=instance.id,
        candidate_id=instance.candidate_id,
        candidate_name=candidate.name if candidate else "Candidate",
        candidate_email=candidate.email if candidate else None,
        role_title=role_title,
        current_state=instance.current_state,
        expires_at=instance.token_expires_at,
        candidate_portal_url=portal_url,
        missing_from_company=missing_company,
        missing_from_candidate=missing_candidate,
        required_documents=req_docs,
        uploaded_evidence=uploaded_evidence,
        company_fields=ctx.get("company_fields", {}),
        progress_percentage=progress,
        created_at=instance.created_at.isoformat() if instance.created_at else None,
    )


@router.get("/instances", response_model=List[OnboardingInstanceSummary])
def list_workflow_instances(
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Lists all running and completed candidate onboarding instances for the tenant."""
    instances = db.query(HRWorkflowInstance).filter_by(
        tenant_id=current_user.tenant_id,
    ).order_by(HRWorkflowInstance.created_at.desc()).all()

    return [_serialize_instance_summary(db, inst) for inst in instances]


@router.post("/instances", response_model=OnboardingInstanceSummary)
def create_onboarding_instance(
    payload: OnboardingInstanceCreateRequest,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Launches a new candidate onboarding instance with custom document requirements and company fields."""
    candidate = db.query(Candidate).filter_by(
        id=payload.candidate_id,
        tenant_id=current_user.tenant_id,
    ).first()
    if not candidate:
        raise HTTPException(status_code=404, detail="Candidate not found")

    wf_def_id = payload.workflow_definition_id
    if not wf_def_id:
        default_def = db.query(HRWorkflowDefinition).filter_by(
            tenant_id=current_user.tenant_id,
            is_active=True,
        ).first()
        if not default_def:
            default_def = HRWorkflowDefinition(
                tenant_id=current_user.tenant_id,
                name="Standard Candidate Onboarding",
                graph_schema={"nodes": [], "edges": []},
            )
            db.add(default_def)
            db.flush()
        wf_def_id = default_def.id

    app_id = payload.application_id
    if not app_id:
        existing_app = db.query(Application).filter_by(
            candidate_id=candidate.id,
            tenant_id=current_user.tenant_id,
        ).first()
        if existing_app:
            app_id = existing_app.id
        else:
            new_app = Application(
                tenant_id=current_user.tenant_id,
                candidate_id=candidate.id,
                status="offered",
                current_stage="offered",
            )
            db.add(new_app)
            db.flush()
            app_id = new_app.id

    raw_token, token_hash, expires_at = CandidatePortalService.generate_token_pair(expires_in_days=7)

    context = dict(payload.context_data or {})
    context["company_fields"] = payload.company_fields
    context["raw_portal_token"] = raw_token
    if payload.required_documents:
        context["required_documents"] = [doc.model_dump() for doc in payload.required_documents]
    elif "required_documents" not in context:
        # Deduce requirements directly from the workflow's linked templates
        deduced_pool = []
        if wf_def_id:
            wf_def = db.query(HRWorkflowDefinition).filter_by(id=wf_def_id).first()
            if wf_def and wf_def.graph_schema:
                nodes = wf_def.graph_schema.get("nodes", [])
                template_ids = [n.get("config", {}).get("template_id") for n in nodes if n.get("config", {}).get("template_id")]
                if template_ids:
                    tpl_versions = db.query(TemplateVersion).filter(
                        TemplateVersion.template_id.in_(template_ids),
                        TemplateVersion.tenant_id == current_user.tenant_id,
                    ).all()
                    for tv in tpl_versions:
                        meta = tv.docx_structure or {}
                        if "deduced_requirements" in meta and isinstance(meta["deduced_requirements"], list):
                            deduced_pool.extend(meta["deduced_requirements"])

        if deduced_pool:
            seen_cats = set()
            unique_deduced = []
            for d in deduced_pool:
                cat = d.get("document_category")
                if cat and cat not in seen_cats:
                    seen_cats.add(cat)
                    unique_deduced.append(d)
            context["required_documents"] = unique_deduced
        else:
            context["required_documents"] = [
                doc.model_dump() for doc in CandidatePortalService.get_required_documents(context)
            ]

    instance = HRWorkflowInstance(
        tenant_id=current_user.tenant_id,
        workflow_definition_id=wf_def_id,
        application_id=app_id,
        candidate_id=candidate.id,
        current_state="intake_pending",
        candidate_token_hash=token_hash,
        token_expires_at=expires_at,
        context_data=context,
    )
    db.add(instance)
    db.flush()

    if candidate.name:
        db.add(CandidateFieldValue(
            tenant_id=current_user.tenant_id,
            instance_id=instance.id,
            canonical_key="candidate.full_name",
            candidate_provided_value=candidate.name,
            resolved_value=candidate.name,
            status="reconciled",
        ))
    if candidate.email:
        db.add(CandidateFieldValue(
            tenant_id=current_user.tenant_id,
            instance_id=instance.id,
            canonical_key="candidate.email",
            candidate_provided_value=candidate.email,
            resolved_value=candidate.email,
            status="reconciled",
        ))
    if candidate.phone:
        db.add(CandidateFieldValue(
            tenant_id=current_user.tenant_id,
            instance_id=instance.id,
            canonical_key="candidate.phone",
            candidate_provided_value=candidate.phone,
            resolved_value=candidate.phone,
            status="reconciled",
        ))

    db.commit()
    db.refresh(instance)
    return _serialize_instance_summary(db, instance)


@router.patch("/instances/{instance_id}", response_model=OnboardingInstanceSummary)
def update_onboarding_instance(
    instance_id: str,
    payload: OnboardingInstanceUpdateRequest,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Updates company fields, document requirements, or status for an onboarding instance."""
    instance = db.query(HRWorkflowInstance).filter_by(
        id=instance_id,
        tenant_id=current_user.tenant_id,
    ).first()
    if not instance:
        raise HTTPException(status_code=404, detail="Instance not found")

    context = dict(instance.context_data or {})
    if payload.company_fields is not None:
        existing_cf = context.get("company_fields") or {}
        existing_cf.update(payload.company_fields)
        context["company_fields"] = existing_cf

    if payload.required_documents is not None:
        context["required_documents"] = [doc.model_dump() for doc in payload.required_documents]

    if payload.context_data is not None:
        context.update(payload.context_data)

    instance.context_data = context

    if payload.current_state:
        instance.current_state = payload.current_state

    db.commit()
    db.refresh(instance)
    return _serialize_instance_summary(db, instance)


@router.get("/instances/{instance_id}")
def get_instance_details(
    instance_id: str,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Retrieves full workflow instance state, tasks, and candidate progress."""
    instance = db.query(HRWorkflowInstance).filter_by(
        id=instance_id,
        tenant_id=current_user.tenant_id,
    ).first()
    if not instance:
        raise HTTPException(status_code=404, detail="Instance not found")

    tasks = db.query(HRWorkflowTask).filter_by(instance_id=instance.id).all()
    evidence = db.query(CandidateEvidence).filter_by(instance_id=instance.id).all()
    fields = db.query(CandidateFieldValue).filter_by(instance_id=instance.id).all()

    return {
        "id": instance.id,
        "workflow_definition_id": instance.workflow_definition_id,
        "candidate_id": instance.candidate_id,
        "current_state": instance.current_state,
        "expires_at": instance.token_expires_at.isoformat(),
        "tasks": [
            {"id": t.id, "node_id": t.node_id, "type": t.task_type, "status": t.status}
            for t in tasks
        ],
        "evidence_count": len(evidence),
        "fields": {f.canonical_key: f.resolved_value for f in fields},
    }


@router.post("/instances/{instance_id}/reconcile", response_model=ReconciliationSummary)
def reconcile_instance(
    instance_id: str,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Reconciles candidate provided data with OCR evidence and flags conflicts."""
    instance = db.query(HRWorkflowInstance).filter_by(
        id=instance_id,
        tenant_id=current_user.tenant_id,
    ).first()
    if not instance:
        raise HTTPException(status_code=404, detail="Instance not found")

    summary = EvidenceReconciliationService.reconcile_instance_data(db, instance_id)
    if not summary.is_blocked:
        instance.current_state = "review_ready"
        db.commit()
    return summary


# --- REVIEW STUDIO ENDPOINTS ---

@router.get("/drafts/{draft_id}/preview", response_model=DraftPreviewResponse)
def get_draft_preview(
    draft_id: str,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Fetches high-fidelity paginated preview pages for Review Studio."""
    draft = db.query(DocumentDraft).filter_by(
        id=draft_id,
        tenant_id=current_user.tenant_id,
    ).first()
    if not draft:
        raise HTTPException(status_code=404, detail="Document draft not found")

    # Load content
    content_bytes = DocumentStore.read(draft.storage_key)
    pages = ReviewStudioRenderer.generate_paginated_preview(content_bytes)

    # Fetch conflicts if any
    conflicts_recs = db.query(CandidateFieldValue).filter_by(
        instance_id=draft.instance_id,
        has_conflict=True,
    ).all()
    conflicts = [
        {
            "canonical_key": c.canonical_key,
            "candidate_value": c.candidate_provided_value,
            "extracted_value": c.document_extracted_value,
            "source_snippet": (c.source_provenance or {}).get("snippet"),
            "confidence": (c.source_provenance or {}).get("confidence", 0.0),
            "reason": (c.conflict_details or {}).get("reason", "Conflict detected"),
        }
        for c in conflicts_recs
    ]

    comments = db.query(DocumentReviewComment).filter_by(draft_id=draft.id).all()

    return DraftPreviewResponse(
        draft_id=draft.id,
        document_name=draft.document_name,
        template_version_id=draft.template_version_id,
        current_approval_status=draft.current_approval_status,
        content_sha256=draft.content_sha256,
        field_values=draft.field_overrides or {},
        pages=pages,
        conflicts=conflicts,
        revisions=draft.editor_revisions or [],
        comments=[
            {"id": c.id, "user_name": c.user_name, "paragraph_index": c.paragraph_index, "text": c.comment_text}
            for c in comments
        ],
    )


@router.put("/drafts/{draft_id}/edit")
def edit_draft(
    draft_id: str,
    req: DraftEditRequest,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Edits draft content and strictly invalidates prior approvals."""
    draft = ReviewStudioRenderer.apply_draft_edit(
        db=db,
        draft_id=draft_id,
        user_id=current_user.id,
        field_overrides=req.field_overrides,
        paragraph_updates=req.paragraph_updates,
        edit_summary=req.edit_summary,
    )
    return {
        "status": "updated",
        "current_approval_status": draft.current_approval_status,
        "approval_invalidated": draft.current_approval_status == "draft",
    }


@router.post("/drafts/{draft_id}/approve")
def approve_draft(
    draft_id: str,
    req: DraftApprovalRequest,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Authorizes document draft, binding cryptographically to content SHA-256."""
    draft = db.query(DocumentDraft).filter_by(
        id=draft_id,
        tenant_id=current_user.tenant_id,
    ).first()
    if not draft:
        raise HTTPException(status_code=404, detail="Draft not found")

    draft.current_approval_status = "approved"
    draft.approved_by_user_id = current_user.id
    draft.approved_at = datetime.now(timezone.utc)
    draft.bound_approval_sha256 = draft.content_sha256
    db.commit()

    return {
        "status": "approved",
        "bound_approval_sha256": draft.bound_approval_sha256,
        "approved_at": draft.approved_at.isoformat(),
    }


@router.post("/drafts/{draft_id}/schedule")
def schedule_draft_delivery(
    draft_id: str,
    req: DraftScheduleRequest,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Schedules dispatch of an approved document in the configured timezone."""
    try:
        delivery = DeliverySchedulerService.schedule_delivery(
            db=db,
            draft_id=draft_id,
            scheduled_for=req.scheduled_for,
            timezone_str=req.timezone,
            recipient_email=req.recipient_email,
        )
        return {
            "delivery_id": delivery.id,
            "status": delivery.status,
            "scheduled_for": delivery.scheduled_for.isoformat(),
            "timezone": delivery.timezone,
        }
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.post("/deliveries/{delivery_id}/dispatch")
def trigger_dispatch(
    delivery_id: str,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Executes pre-flight revalidation and dispatches the document."""
    success, message = DeliverySchedulerService.revalidate_and_dispatch(db, delivery_id)
    if not success:
        raise HTTPException(status_code=409, detail=message)
    return {"status": "dispatched", "message": message}


@router.post("/instances/{instance_id}/verify-signed-document")
def verify_signed_document(
    instance_id: str,
    evidence_id: str,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Executes multi-dimensional verification on candidate signed evidence."""
    try:
        report = DocumentVerificationService.verify_instance_signed_evidence(
            db=db,
            tenant_id=current_user.tenant_id,
            instance_id=instance_id,
            evidence_id=evidence_id,
        )
        return report
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))


@router.post("/instances/{instance_id}/hr-verify-signature")
def hr_verify_signature(
    instance_id: str,
    evidence_id: str,
    notes: Optional[str] = None,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Records mandatory HR director verification sign-off on executed document."""
    try:
        result = DocumentVerificationService.record_hr_verification_signoff(
            db=db,
            tenant_id=current_user.tenant_id,
            instance_id=instance_id,
            evidence_id=evidence_id,
            hr_user_id=current_user.id,
            notes=notes,
        )
        return result
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))


@router.post("/instances/{instance_id}/complete-onboarding-and-provision")
def complete_onboarding_and_provision(
    instance_id: str,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_user),
):
    """Completes onboarding workflow and atomically provisions active employee department-wise."""
    try:
        result = EmployeeProvisioningService.provision_employee_from_onboarding(
            db=db,
            tenant_id=current_user.tenant_id,
            instance_id=instance_id,
            actor_id=current_user.id,
        )
        return result
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
