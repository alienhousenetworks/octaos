from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import List, Any, Optional, Dict
from pydantic import BaseModel
from datetime import datetime
from pathlib import Path
import re
import io
import csv
import logging

logger = logging.getLogger(__name__)

from app.core.config import settings

from app.api import deps
from app.core.rbac import Action, Resource, require_permission
from app.models.base import User, APICredential
from app.schemas import verticals as schemas
from app.models import verticals as models
from app.services.agents.hr import HRAgent
from app.services.llm_gateway import LLMGateway
from app.services.hr.resume_parser import (
    extract_resume_text,
    extract_heuristic_entities,
    calculate_ats_keyword_score,
    evaluate_semantic_fit,
    compute_composite_ranking,
)
from app.services.hr.contract_service import (
    get_available_templates,
    render_contract_document,
    generate_offer_email_content,
)

router = APIRouter()

class SourceRequest(BaseModel):
    role: str
    requirements: str
    salary: str
    count: int = 5
    platforms: List[str] = ["linkedin", "indeed", "ziprecruiter", "greenhouse", "lever"]
    provider: str = "auto"
    model: Optional[str] = None

class OutreachRequest(BaseModel):
    channel: str = "free_outreach"
    subject: str = "Exciting job opportunity: {role}"
    body_template: str = "Hello {name},\n\nI saw your profile and thought you would be a great fit for our {role} opening..."
    provider: str = "auto"
    model: Optional[str] = None

class InterviewRequest(BaseModel):
    tool: str = "free_scheduling"
    provider: str = "auto"
    model: Optional[str] = None

class InterviewEmailRequest(BaseModel):
    interview_date: str
    interview_time: Optional[str] = "11:00 AM"
    round_name: Optional[str] = "Technical Screen"
    meeting_link: Optional[str] = "https://meet.google.com/new"
    interviewer_name: Optional[str] = "Engineering Lead"
    custom_notes: Optional[str] = None
    channel: str = "smtp"

class GenerateOfferRequest(BaseModel):
    template_id: str = "full_time_standard"
    variables: Dict[str, Any] = {}

class SendOfferEmailRequest(BaseModel):
    template_id: str = "full_time_standard"
    variables: Dict[str, Any] = {}
    channel: str = "smtp"

class StatusUpdateRequest(BaseModel):
    status: str

@router.get("/", response_model=List[schemas.Candidate])
def read_candidates(
    db: Session = Depends(deps.get_db),
    tenant_id: str = Depends(deps.get_current_tenant_id),
    _: User = Depends(require_permission(Resource.HR, Action.READ)),
    skip: int = 0,
    limit: int = 100
) -> Any:
    candidates = db.query(models.Candidate).filter(
        models.Candidate.tenant_id == tenant_id
    ).order_by(models.Candidate.created_at.desc()).offset(skip).limit(limit).all()
    return candidates

@router.post("/", response_model=schemas.Candidate)
def create_candidate(
    *,
    db: Session = Depends(deps.get_db),
    tenant_id: str = Depends(deps.get_current_tenant_id),
    _: User = Depends(require_permission(Resource.HR, Action.CREATE)),
    candidate_in: schemas.CandidateCreate
) -> Any:
    candidate = models.Candidate(**candidate_in.dict(), tenant_id=tenant_id)
    db.add(candidate)
    db.commit()
    db.refresh(candidate)
    return candidate

def handle_value_error(e: ValueError):
    ve_str = str(e)
    provider = None
    for p in ["linkedin", "meta", "facebook", "instagram", "twitter", "gmail", "whatsapp", "apollo", "hunter", "google_places", "google_calendar", "smtp", "greenhouse", "lever", "openai", "anthropic", "gemini"]:
        if p in ve_str.lower():
            provider = p
            break
    if provider:
        if provider == "smtp":
            msg = "I need your SMTP outgoing mail credentials. Please reply with: 'My smtp credential is: smtp://username:password@smtp.mailtrap.io:2525'."
        else:
            msg = f"I need your {provider} API key to complete this task. Please reply with 'My {provider} key is: [YOUR_KEY]'."
        return {"status": "action_required", "message": msg}
    raise HTTPException(status_code=400, detail=str(e))

@router.post("/source")
async def source_candidates(
    *,
    db: Session = Depends(deps.get_db),
    tenant_id: str = Depends(deps.get_current_tenant_id),
    _: User = Depends(require_permission(Resource.HR, Action.CREATE)),
    request: SourceRequest
) -> Any:
    agent = HRAgent(db, tenant_id)
    try:
        result = await agent.execute_task({
            "action": "source_candidates",
            "parameters": {
                "role": request.role,
                "requirements": request.requirements,
                "salary": request.salary,
                "count": request.count,
                "platforms": request.platforms,
                "provider": request.provider,
                "model": request.model
            }
        })
        return result
    except ValueError as e:
        return handle_value_error(e)

@router.post("/{candidate_id}/outreach")
async def candidate_outreach(
    candidate_id: str,
    *,
    db: Session = Depends(deps.get_db),
    tenant_id: str = Depends(deps.get_current_tenant_id),
    _: User = Depends(require_permission(Resource.HR, Action.CREATE)),
    request: OutreachRequest
) -> Any:
    agent = HRAgent(db, tenant_id)
    try:
        result = await agent.execute_task({
            "action": "candidate_outreach",
            "parameters": {
                "candidate_id": candidate_id,
                "channel": request.channel,
                "subject": request.subject,
                "body_template": request.body_template,
                "provider": request.provider,
                "model": request.model
            }
        })
        return result
    except ValueError as e:
        return handle_value_error(e)

@router.post("/{candidate_id}/interview")
async def schedule_interview(
    candidate_id: str,
    *,
    db: Session = Depends(deps.get_db),
    tenant_id: str = Depends(deps.get_current_tenant_id),
    _: User = Depends(require_permission(Resource.HR, Action.CREATE)),
    request: InterviewRequest
) -> Any:
    agent = HRAgent(db, tenant_id)
    try:
        result = await agent.execute_task({
            "action": "schedule_interview",
            "parameters": {
                "candidate_id": candidate_id,
                "tool": request.tool,
                "provider": request.provider,
                "model": request.model
            }
        })
        return result
    except ValueError as e:
        return handle_value_error(e)

@router.post("/{candidate_id}/status", response_model=schemas.Candidate)
@router.put("/{candidate_id}", response_model=schemas.Candidate)
def update_candidate(
    candidate_id: str,
    *,
    db: Session = Depends(deps.get_db),
    tenant_id: str = Depends(deps.get_current_tenant_id),
    _: User = Depends(require_permission(Resource.HR, Action.CREATE)),
    candidate_in: schemas.CandidateUpdate
) -> Any:
    """Updates candidate properties and scorecard fields from the Extraction Sheet."""
    candidate = db.query(models.Candidate).filter(
        models.Candidate.id == candidate_id,
        models.Candidate.tenant_id == tenant_id
    ).first()
    if not candidate:
        raise HTTPException(status_code=404, detail="Candidate not found")

    if candidate_in.name is not None:
        candidate.name = candidate_in.name
    if candidate_in.email is not None:
        candidate.email = candidate_in.email
    if candidate_in.role is not None:
        candidate.role = candidate_in.role
    if candidate_in.status is not None:
        candidate.status = candidate_in.status

    scorecard = dict(candidate.scorecard or {})
    if candidate_in.phone is not None:
        scorecard["phone_number"] = candidate_in.phone
        scorecard["phone"] = candidate_in.phone
    if candidate_in.address is not None:
        scorecard["address"] = candidate_in.address
    if candidate_in.skills is not None:
        scorecard["skills"] = candidate_in.skills
    if candidate_in.projects is not None:
        scorecard["projects"] = candidate_in.projects
    if candidate_in.experience_years is not None:
        scorecard["experience_years"] = candidate_in.experience_years
    if candidate_in.experience_summary is not None:
        scorecard["experience_summary"] = candidate_in.experience_summary
    if candidate_in.resume_link is not None:
        scorecard["resume_link"] = candidate_in.resume_link
    if candidate_in.extra_remarks is not None:
        scorecard["extra_remarks"] = candidate_in.extra_remarks
    if candidate_in.scorecard is not None:
        scorecard.update(candidate_in.scorecard)

    candidate.scorecard = scorecard
    db.commit()
    db.refresh(candidate)
    return candidate

@router.delete("/{candidate_id}")
def delete_candidate(
    candidate_id: str,
    *,
    db: Session = Depends(deps.get_db),
    tenant_id: str = Depends(deps.get_current_tenant_id),
    _: User = Depends(require_permission(Resource.HR, Action.DELETE)),
) -> Any:
    candidate = db.query(models.Candidate).filter(
        models.Candidate.id == candidate_id,
        models.Candidate.tenant_id == tenant_id
    ).first()
    if not candidate:
        raise HTTPException(status_code=404, detail="Candidate not found")
    
    db.delete(candidate)
    db.commit()
    return {"status": "success"}

class BatchDeleteRequest(BaseModel):
    candidate_ids: List[str]

@router.post("/candidates/batch-delete")
def batch_delete_candidates(
    request: BatchDeleteRequest,
    *,
    db: Session = Depends(deps.get_db),
    tenant_id: str = Depends(deps.get_current_tenant_id),
    _: User = Depends(require_permission(Resource.HR, Action.DELETE)),
) -> Any:
    """Deletes multiple candidates in batch or all matching candidate IDs."""
    if not request.candidate_ids:
        raise HTTPException(status_code=400, detail="No candidate IDs provided for deletion.")
    
    deleted = db.query(models.Candidate).filter(
        models.Candidate.tenant_id == tenant_id,
        models.Candidate.id.in_(request.candidate_ids)
    ).delete(synchronize_session=False)
    
    db.commit()
    return {
        "status": "success",
        "message": f"Successfully deleted {deleted} candidate(s).",
        "deleted_count": deleted
    }

# --- BATCH RESUME UPLOAD, ATS MATCHING & RANKING ---

@router.get("/resumes/ingestion-status")
def get_resume_ingestion_status(
    db: Session = Depends(deps.get_db),
    tenant_id: str = Depends(deps.get_current_tenant_id),
    _: User = Depends(require_permission(Resource.HR, Action.READ)),
) -> Any:
    # Ensure table exists checkfirst
    try:
        models.HRResumeIngestionJob.__table__.create(bind=db.get_bind(), checkfirst=True)
    except Exception as create_err:
        logger.warning(f"Failed to check/create hr_resume_ingestion_jobs table: {create_err}")

    try:
        job = db.query(models.HRResumeIngestionJob).filter(
            models.HRResumeIngestionJob.tenant_id == tenant_id
        ).order_by(models.HRResumeIngestionJob.created_at.desc()).first()
    except Exception as q_err:
        logger.warning(f"Error querying HRResumeIngestionJob: {q_err}")
        return {"has_active_job": False, "job": None}

    if not job:
        return {"has_active_job": False, "job": None}

    return {
        "has_active_job": job.status == "running",
        "job": {
            "id": job.id,
            "role": job.role,
            "total_files": job.total_files,
            "processed_files": job.processed_files,
            "progress_percent": job.progress_percent,
            "current_file": job.current_file,
            "status": job.status,
            "logs": job.logs or [],
            "error_message": job.error_message,
            "created_at": job.created_at.isoformat() if job.created_at else None,
            "updated_at": job.updated_at.isoformat() if job.updated_at else None,
        }
    }

@router.post("/resumes/upload-batch")
async def upload_resumes_batch(
    files: List[UploadFile] = File(...),
    role: str = Form(...),
    requirements: str = Form(...),
    salary: str = Form("$120,000/year"),
    provider: str = Form("auto"),
    model: Optional[str] = Form(None),
    db: Session = Depends(deps.get_db),
    tenant_id: str = Depends(deps.get_current_tenant_id),
    _: User = Depends(require_permission(Resource.HR, Action.CREATE)),
) -> Any:
    """Accepts multiple resume files or folder uploads, extracts text, computes zero-cost ATS match,

    performs condensed semantic fit evaluation, and ranks all candidates.
    """
    if not files:
        raise HTTPException(status_code=400, detail="No resume files uploaded.")

    llm_gateway = LLMGateway(db, tenant_id)
    raw_candidates = []
    total_files = len(files)
    logger.info(f"[HR Ingestion] Starting batch processing for {total_files} resume(s) | Role: {role}")

    # Ensure table exists checkfirst
    try:
        models.HRResumeIngestionJob.__table__.create(bind=db.get_bind(), checkfirst=True)
    except Exception as create_err:
        logger.warning(f"Failed to check/create hr_resume_ingestion_jobs table: {create_err}")

    # Track persistent job in database
    job = models.HRResumeIngestionJob(
        tenant_id=tenant_id,
        role=role,
        total_files=total_files,
        processed_files=0,
        progress_percent=5,
        current_file=f"Preparing {total_files} resume(s)...",
        status="running",
        logs=[
            f"[{datetime.now().strftime('%H:%M:%S')}] Started batch ingestion of {total_files} resumes for role: {role}",
            f"[{datetime.now().strftime('%H:%M:%S')}] Using engine: {provider.upper()} ({model or 'default'})"
        ]
    )
    db.add(job)
    db.commit()
    db.refresh(job)

    # Ensure tenant resume media directory exists
    resumes_dir = Path(settings.MEDIA_UPLOAD_DIR) / "resumes" / tenant_id
    resumes_dir.mkdir(parents=True, exist_ok=True)

    try:
        for idx, file in enumerate(files, start=1):
            # Read file bytes
            content = await file.read()
            raw_filename = file.filename or f"resume_{idx}.pdf"
            clean_fn = re.sub(r"[^a-zA-Z0-9_.-]", "_", raw_filename)
            timestamp_prefix = datetime.now().strftime("%Y%m%d%H%M%S")
            safe_filename = f"{timestamp_prefix}_{clean_fn}"

            calc_percent = int(10 + ((idx - 1) / total_files) * 75)
            job.current_file = raw_filename
            job.processed_files = idx - 1
            job.progress_percent = calc_percent
            job.logs = [*(job.logs or []), f"[{datetime.now().strftime('%H:%M:%S')}] ({idx}/{total_files}) Parsing '{raw_filename}'..."]
            db.commit()

            logger.info(f"[HR Ingestion] ({idx}/{total_files}) Ingesting '{raw_filename}' ({len(content)} bytes)...")

            # Persist uploaded resume locally for internal OctaOS access
            file_disk_path = resumes_dir / safe_filename
            try:
                with open(file_disk_path, "wb") as f_out:
                    f_out.write(content)
                resume_link = f"/media/resumes/{tenant_id}/{safe_filename}"
            except Exception as save_err:
                logger.warning(f"[HR Ingestion] Could not save resume disk copy for '{raw_filename}': {save_err}")
                resume_link = ""

            # Tier 1: Deterministic Local Text Extraction ($0 cost)
            extracted_text = extract_resume_text(raw_filename, content)
            if not extracted_text:
                try:
                    extracted_text = content.decode("utf-8", errors="ignore")
                except Exception:
                    extracted_text = ""

            logger.info(f"[HR Ingestion] ({idx}/{total_files}) Extracted {len(extracted_text)} chars of text from '{raw_filename}'")

            # Tier 2: Local Heuristic Entity & ATS Keyword Match ($0 cost)
            entities = extract_heuristic_entities(extracted_text, filename=raw_filename)
            ats_results = calculate_ats_keyword_score(entities["skills"], extracted_text, requirements, role)

            condensed_profile = {
                **entities,
                **ats_results,
                "filename": raw_filename,
                "resume_link": resume_link,
                "text_excerpt": extracted_text[:1500]
            }

            # Tier 3: Condensed Semantic Evaluation via Low-Cost / Fast LLM (<500 tokens)
            try:
                semantic_res = await evaluate_semantic_fit(
                    llm_gateway=llm_gateway,
                    condensed_profile=condensed_profile,
                    role=role,
                    requirements=requirements,
                    provider=provider,
                    model=model
                )
            except Exception as e:
                rec = "Recommended" if ats_results.get("ats_score", 50) >= 70 else "Consider"
                verdict = f"Assessed via deterministic heuristics. {len(ats_results.get('matched_skills', []))} core skills matched."
                semantic_res = {
                    "semantic_score": ats_results.get("ats_score", 50),
                    "experience_summary": f"Candidate with {entities.get('experience_years', 0)} years experience.",
                    "strengths": [f"Skill match: {s}" for s in ats_results.get("matched_skills", [])[:3]],
                    "gaps": [f"Missing: {s}" for s in ats_results.get("missing_skills", [])[:2]],
                    "recommendation": rec,
                    "fit_verdict": verdict,
                    "extra_remarks": f"{rec}: {verdict}"
                }

            candidate_data = {
                **condensed_profile,
                **semantic_res,
                "filename": raw_filename,
                "resume_link": resume_link,
                "raw_text_excerpt": extracted_text[:800]
            }
            raw_candidates.append(candidate_data)
            logger.info(
                f"[HR Ingestion] ({idx}/{total_files}) Candidate parsed: '{entities.get('name')}' | "
                f"Phone: '{entities.get('phone')}' | ATS Score: {ats_results.get('ats_score')}%"
            )

        # Compute composite rank across the batch
        job.progress_percent = 90
        job.current_file = "Computing composite ranking & writing candidates..."
        db.commit()

        ranked_candidates = compute_composite_ranking(raw_candidates)
        logger.info(f"[HR Ingestion] Successfully ranked {len(ranked_candidates)} candidate profiles. Committing to DB...")

        created_or_updated = []
        for item in ranked_candidates:
            cand_email = item.get("email")
            existing = None
            if cand_email:
                existing = db.query(models.Candidate).filter_by(
                    tenant_id=tenant_id, email=cand_email
                ).first()

            scorecard = {
                "rank": item.get("rank", 1),
                "composite_score": item.get("composite_score", 60),
                "ats_score": item.get("ats_score", 60),
                "semantic_score": item.get("semantic_score", 60),
                "match_score": item.get("composite_score", 60),
                "skills": item.get("skills", []),
                "matched_skills": item.get("matched_skills", []),
                "missing_skills": item.get("missing_skills", []),
                "experience_years": item.get("experience_years", 0),
                "experience_summary": item.get("experience_summary", ""),
                "education": item.get("education", ""),
                "address": item.get("address", "Not specified"),
                "projects": item.get("projects", []),
                "phone_number": item.get("phone", ""),
                "linkedin": item.get("linkedin", ""),
                "github": item.get("github", ""),
                "strengths": item.get("strengths", []),
                "gaps": item.get("gaps", []),
                "recommendation": item.get("recommendation", "Consider"),
                "fit_verdict": item.get("fit_verdict", ""),
                "extra_remarks": item.get("extra_remarks") or f"{item.get('recommendation', 'Consider')}: {item.get('fit_verdict', '')}",
                "requirements_match": f"Matched {len(item.get('matched_skills', []))} required skills. {item.get('fit_verdict', '')}",
                "salary_expectation": salary,
                "filename": item.get("filename", ""),
                "resume_link": item.get("resume_link", ""),
                "source_platform": "Batch Resume Ingestion",
                "resume_excerpt": item.get("raw_text_excerpt", "")
            }

            if existing:
                existing.role = role
                existing.scorecard = scorecard
                existing.status = "sourced"
                created_or_updated.append(existing)
            else:
                new_cand = models.Candidate(
                    tenant_id=tenant_id,
                    name=item.get("name") if item.get("name") and item.get("name") != "Candidate" else "N/A",
                    email=cand_email if cand_email else "N/A",
                    role=role,
                    scorecard=scorecard,
                    status="sourced"
                )
                db.add(new_cand)
                created_or_updated.append(new_cand)

        job.status = "completed"
        job.processed_files = total_files
        job.progress_percent = 100
        job.current_file = "Completed successfully"
        job.logs = [*(job.logs or []), f"[{datetime.now().strftime('%H:%M:%S')}] ✓ Ingestion completed. Saved {len(created_or_updated)} candidates to database."]
        db.commit()

        for c in created_or_updated:
            db.refresh(c)

        return {
            "status": "success",
            "job_id": job.id,
            "processed_count": len(ranked_candidates),
            "candidates": [schemas.Candidate.from_orm(c) for c in created_or_updated]
        }
    except Exception as e:
        logger.error(f"[HR Ingestion] Batch processing failed: {e}", exc_info=True)
        job.status = "failed"
        job.error_message = str(e)
        job.logs = [*(job.logs or []), f"[{datetime.now().strftime('%H:%M:%S')}] ✗ Failed: {str(e)}"]
        db.commit()
        raise HTTPException(status_code=500, detail=f"Resume ingestion failed: {str(e)}")


# --- INTERVIEW INVITATION EMAIL ---

@router.post("/{candidate_id}/send-interview-email")
async def send_interview_email(
    candidate_id: str,
    request: InterviewEmailRequest,
    db: Session = Depends(deps.get_db),
    tenant_id: str = Depends(deps.get_current_tenant_id),
    _: User = Depends(require_permission(Resource.HR, Action.CREATE)),
) -> Any:
    """Dispatches a personalized interview invitation email to the approved candidate."""
    candidate = db.query(models.Candidate).filter(
        models.Candidate.id == candidate_id,
        models.Candidate.tenant_id == tenant_id
    ).first()
    if not candidate:
        raise HTTPException(status_code=404, detail="Candidate not found")

    company_name = "OctaOS Technologies Inc."
    subject = f"Interview Invitation: {candidate.role} at {company_name}"
    
    body = f"""Dear {candidate.name},

Thank you for your application for the position of {candidate.role} at {company_name}. 
We reviewed your profile and were impressed by your background. We would love to invite you to an interview!

Interview Details:
- Round: {request.round_name}
- Date: {request.interview_date}
- Time: {request.interview_time}
- Interviewer: {request.interviewer_name}
- Meeting Link: {request.meeting_link}

{f"Additional Notes:\n{request.custom_notes}\n" if request.custom_notes else ""}
Please reply directly to this email if you need to reschedule or have any questions.

We look forward to speaking with you!

Warm regards,
Talent Acquisition Team
{company_name}
"""

    agent = HRAgent(db, tenant_id)
    sent_successfully = False
    delivery_note = "Logged interview invite"

    # Attempt SMTP or Gmail dispatch if credentials exist
    try:
        if request.channel == "smtp":
            cred = db.query(APICredential).filter_by(tenant_id=tenant_id, provider="smtp").first()
            if cred and cred.encrypted_key:
                from app.core.security import decrypt_api_key
                smtp_creds = agent._parse_smtp_credentials(decrypt_api_key(cred.encrypted_key))
                if smtp_creds and candidate.email and "@" in candidate.email:
                    sent_successfully = agent._send_actual_email(smtp_creds, candidate.email, subject, body)
                    delivery_note = "Sent via SMTP"
        elif request.channel == "gmail":
            cred = db.query(APICredential).filter_by(tenant_id=tenant_id, provider="gmail").first()
            if cred and cred.encrypted_key:
                from app.core.security import decrypt_api_key
                agent._send_gmail_api_email(decrypt_api_key(cred.encrypted_key), candidate.email, subject, body)
                sent_successfully = True
                delivery_note = "Sent via Gmail API"
    except Exception as e:
        delivery_note = f"Delivery attempted (Note: {str(e)[:100]})"

    # Update candidate status and scorecard
    candidate.status = "interviewed"
    existing_scorecard = candidate.scorecard or {}
    candidate.scorecard = {
        **existing_scorecard,
        "meeting_time": f"{request.interview_date} at {request.interview_time}",
        "meeting_link": request.meeting_link,
        "interview_round": request.round_name,
        "interviewer_name": request.interviewer_name,
        "interview_email_subject": subject,
        "interview_email_body": body,
        "interview_delivery_status": delivery_note,
        "calendar_booked": True
    }
    db.commit()
    db.refresh(candidate)

    return {
        "status": "success",
        "message": f"Interview invitation successfully recorded. ({delivery_note})",
        "candidate": schemas.Candidate.from_orm(candidate)
    }


# --- OFFER LETTER & CONTRACT TEMPLATES STUDIO ---

@router.get("/contract-templates")
def list_contract_templates(
    _: User = Depends(require_permission(Resource.HR, Action.READ)),
) -> Any:
    """Returns available contract templates and default attributes."""
    return get_available_templates()


@router.post("/{candidate_id}/generate-offer")
def generate_offer_document(
    candidate_id: str,
    request: GenerateOfferRequest,
    db: Session = Depends(deps.get_db),
    tenant_id: str = Depends(deps.get_current_tenant_id),
    _: User = Depends(require_permission(Resource.HR, Action.READ)),
) -> Any:
    """Renders an employment contract or offer letter with dynamic variable substitution."""
    candidate = db.query(models.Candidate).filter(
        models.Candidate.id == candidate_id,
        models.Candidate.tenant_id == tenant_id
    ).first()
    if not candidate:
        raise HTTPException(status_code=404, detail="Candidate not found")

    vars_payload = {
        "candidate_name": candidate.name,
        "candidate_email": candidate.email,
        "job_title": candidate.role,
        **(request.variables or {})
    }

    doc = render_contract_document(request.template_id, vars_payload)
    return {
        "status": "success",
        "document": doc
    }


@router.post("/{candidate_id}/send-offer-email")
async def send_offer_email(
    candidate_id: str,
    request: SendOfferEmailRequest,
    db: Session = Depends(deps.get_db),
    tenant_id: str = Depends(deps.get_current_tenant_id),
    _: User = Depends(require_permission(Resource.HR, Action.CREATE)),
) -> Any:
    """Generates official offer contract and dispatches offer email to candidate."""
    candidate = db.query(models.Candidate).filter(
        models.Candidate.id == candidate_id,
        models.Candidate.tenant_id == tenant_id
    ).first()
    if not candidate:
        raise HTTPException(status_code=404, detail="Candidate not found")

    vars_payload = {
        "candidate_name": candidate.name,
        "candidate_email": candidate.email,
        "job_title": candidate.role,
        **(request.variables or {})
    }

    # Render contract
    doc = render_contract_document(request.template_id, vars_payload)
    company_name = vars_payload.get("company_name", "OctaOS Technologies Inc.")
    start_date = vars_payload.get("start_date", "Immediate")
    salary_amount = vars_payload.get("salary_amount", "120,000")
    salary_currency = vars_payload.get("salary_currency", "USD")

    subject, body = generate_offer_email_content(
        candidate.name, candidate.role, company_name, start_date, salary_amount, salary_currency
    )

    agent = HRAgent(db, tenant_id)
    sent_successfully = False
    delivery_note = "Offer draft registered in system"

    try:
        if request.channel == "smtp":
            cred = db.query(APICredential).filter_by(tenant_id=tenant_id, provider="smtp").first()
            if cred and cred.encrypted_key:
                from app.core.security import decrypt_api_key
                smtp_creds = agent._parse_smtp_credentials(decrypt_api_key(cred.encrypted_key))
                if smtp_creds and candidate.email and "@" in candidate.email:
                    sent_successfully = agent._send_actual_email(smtp_creds, candidate.email, subject, body)
                    delivery_note = "Sent via SMTP"
        elif request.channel == "gmail":
            cred = db.query(APICredential).filter_by(tenant_id=tenant_id, provider="gmail").first()
            if cred and cred.encrypted_key:
                from app.core.security import decrypt_api_key
                agent._send_gmail_api_email(decrypt_api_key(cred.encrypted_key), candidate.email, subject, body)
                sent_successfully = True
                delivery_note = "Sent via Gmail API"
    except Exception as e:
        delivery_note = f"Delivery registered ({str(e)[:100]})"

    # Update candidate status to offered
    candidate.status = "offered"
    existing_scorecard = candidate.scorecard or {}
    candidate.scorecard = {
        **existing_scorecard,
        "offer_details": {
            "template_id": request.template_id,
            "template_name": doc["template_name"],
            "variables": vars_payload,
            "email_subject": subject,
            "email_body": body,
            "contract_html": doc["html"],
            "delivery_note": delivery_note,
            "offered_at": datetime.now().isoformat()
        }
    }
    db.commit()
    db.refresh(candidate)

    return {
        "status": "success",
        "message": f"Offer letter & contract successfully prepared for {candidate.name}. ({delivery_note})",
        "candidate": schemas.Candidate.from_orm(candidate),
        "contract": doc
    }


# --- EXCEL (.XLSX) EXPORT & IMPORT FOR HR CANDIDATES ---

@router.get("/candidates/export-excel")
def export_candidates_excel(
    db: Session = Depends(deps.get_db),
    tenant_id: str = Depends(deps.get_current_tenant_id),
    _: User = Depends(require_permission(Resource.HR, Action.READ)),
):
    """Generates and streams a styled Excel (.xlsx) spreadsheet of all shortlisted & extracted candidates."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    def sanitize_val(v: Any) -> Any:
        if isinstance(v, str):
            return ILLEGAL_CHARACTERS_RE.sub("", v)
        return v

    candidates = db.query(models.Candidate).filter(
        models.Candidate.tenant_id == tenant_id
    ).order_by(models.Candidate.created_at.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Shortlisted Candidates"

    # Define styled columns matching the prompt
    headers = [
        "Rank",
        "Name",
        "Mobile No",
        "Email",
        "Address",
        "Skills",
        "Projects",
        "Experiences",
        "Resume Link (Inside OctaOS)",
        "Shortlisting Status",
        "Match Score (%)",
        "Extra Remarks"
    ]

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")  # Sleek dark gray
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    border_thin = Side(border_style="thin", color="E5E7EB")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    ws.append(headers)

    # Style header row
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = cell_border
    ws.row_dimensions[1].height = 28

    # Populate candidate rows
    for row_idx, cand in enumerate(candidates, start=2):
        scorecard = cand.scorecard or {}
        rank = scorecard.get("rank") or (row_idx - 1)
        name = cand.name if cand.name and cand.name != "Candidate" else "N/A"
        mobile = scorecard.get("phone_number") or scorecard.get("phone") or "N/A"
        raw_email = cand.email or ""
        email = raw_email if raw_email and raw_email != "N/A" and not raw_email.endswith("@applicant.internal") else "N/A"
        address = scorecard.get("address") or "N/A"

        # Skills list
        skills_list = scorecard.get("skills") or []
        skills_str = ", ".join(skills_list) if isinstance(skills_list, list) else str(skills_list)

        # Projects list
        proj_list = scorecard.get("projects") or []
        if isinstance(proj_list, list):
            proj_str = " | ".join(proj_list) if proj_list else "None listed"
        else:
            proj_str = str(proj_list)

        # Experiences
        exp_years = scorecard.get("experience_years", 0)
        exp_summary = scorecard.get("experience_summary") or ""
        experiences_str = f"{exp_years} yrs" + (f" - {exp_summary}" if exp_summary else "")

        # OctaOS resume link
        resume_link = scorecard.get("resume_link") or ""
        if resume_link and not resume_link.startswith("http"):
            resume_link = f"{settings.PUBLIC_BASE_URL.rstrip('/')}{resume_link}"

        status = (cand.status or "sourced").capitalize()
        match_score = scorecard.get("composite_score") or scorecard.get("match_score") or 60
        extra_remarks = scorecard.get("extra_remarks") or scorecard.get("fit_verdict") or scorecard.get("recommendation") or "Shortlisted profile."

        row_values = [
            rank,
            name,
            mobile,
            email,
            address,
            skills_str,
            proj_str,
            experiences_str,
            resume_link,
            status,
            f"{match_score}%",
            extra_remarks
        ]
        clean_row_values = [sanitize_val(v) for v in row_values]
        ws.append(clean_row_values)

        # Format row styling
        ws.row_dimensions[row_idx].height = 24
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            c.border = cell_border
            c.font = Font(name="Calibri", size=10)
            if col_idx in [1, 3, 10, 11]:
                c.alignment = center_align
            else:
                c.alignment = left_align

            # If it's a resume link, make it a clickable hyperlink in Excel
            if col_idx == 9 and resume_link:
                c.hyperlink = resume_link
                c.font = Font(name="Calibri", size=10, color="2563EB", underline="single")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"octaos_candidates_extraction_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.post("/candidates/import-excel")
async def import_candidates_excel(
    file: UploadFile = File(...),
    db: Session = Depends(deps.get_db),
    tenant_id: str = Depends(deps.get_current_tenant_id),
    _: User = Depends(require_permission(Resource.HR, Action.CREATE)),
) -> Any:
    """Accepts an Excel (.xlsx) or CSV spreadsheet and batch imports/updates candidate extraction records."""
    if not file:
        raise HTTPException(status_code=400, detail="No file provided")

    filename = (file.filename or "").lower()
    content = await file.read()

    rows_data: List[Dict[str, Any]] = []

    # 1. Parse Excel XLSX
    if filename.endswith(".xlsx"):
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows or len(rows) < 2:
                raise HTTPException(status_code=400, detail="Excel sheet is empty or missing data rows.")

            raw_headers = [str(h).strip().lower() if h else f"col_{i}" for i, h in enumerate(rows[0])]

            for row_vals in rows[1:]:
                if not any(row_vals):
                    continue
                row_dict = {}
                for idx, col_name in enumerate(raw_headers):
                    val = row_vals[idx] if idx < len(row_vals) else None
                    row_dict[col_name] = str(val).strip() if val is not None else ""
                rows_data.append(row_dict)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to read Excel workbook: {str(e)}")

    # 2. Parse CSV
    elif filename.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig", errors="ignore")
            reader = csv.DictReader(io.StringIO(text))
            for r in reader:
                clean_row = {k.strip().lower(): str(v).strip() for k, v in r.items() if k}
                if any(clean_row.values()):
                    rows_data.append(clean_row)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to read CSV: {str(e)}")
    else:
        raise HTTPException(status_code=400, detail="Unsupported file format. Please upload a .xlsx or .csv file.")

    if not rows_data:
        raise HTTPException(status_code=400, detail="No valid candidate rows found in the uploaded sheet.")

    imported_count = 0
    updated_count = 0

    def find_field(row: Dict[str, str], aliases: List[str]) -> str:
        for a in aliases:
            for k in row:
                if a in k:
                    return row[k]
        return ""

    for row in rows_data:
        name = find_field(row, ["name", "candidate name", "full name"])
        email = find_field(row, ["email", "e-mail", "mail"])
        mobile = find_field(row, ["mobile", "phone", "contact", "cell"])
        address = find_field(row, ["address", "location", "city", "residence"])
        skills_raw = find_field(row, ["skill", "skills", "tech stack"])
        projects_raw = find_field(row, ["project", "projects", "portfolio"])
        exp_raw = find_field(row, ["experience", "experiences", "exp", "years"])
        resume_link = find_field(row, ["resume link", "resume_link", "resume", "cv link"])
        status = find_field(row, ["status", "shortlist", "stage"]).lower() or "sourced"
        remarks = find_field(row, ["remark", "remarks", "extra remark", "extra remarks", "notes", "verdict"])
        score_raw = find_field(row, ["score", "match", "match score", "ats"])

        if not name and not email:
            continue

        # Clean skills into list
        skills = [s.strip() for s in re.split(r"[,|;]", skills_raw) if s.strip()] if skills_raw else []
        projects = [p.strip() for p in re.split(r"[|;\n]", projects_raw) if p.strip()] if projects_raw else []

        # Parse experience years number
        exp_years = 0
        exp_num_match = re.search(r"(\d+)", exp_raw)
        if exp_num_match:
            try:
                exp_years = int(exp_num_match.group(1))
            except:
                exp_years = 0

        # Parse match score
        score = 70
        score_match = re.search(r"(\d+)", score_raw)
        if score_match:
            try:
                score = min(100, max(10, int(score_match.group(1))))
            except:
                score = 70

        # Normalize status
        if status not in ["sourced", "accepted", "screened", "interviewed", "offered", "hired", "rejected"]:
            status = "sourced"

        # Lookup existing candidate by email
        existing = None
        if email and "@" in email:
            existing = db.query(models.Candidate).filter_by(tenant_id=tenant_id, email=email).first()

        scorecard = {
            "rank": len(rows_data),
            "composite_score": score,
            "ats_score": score,
            "semantic_score": score,
            "match_score": score,
            "phone_number": mobile,
            "address": address or "Not specified",
            "skills": skills,
            "projects": projects,
            "experience_years": exp_years,
            "experience_summary": exp_raw,
            "resume_link": resume_link,
            "extra_remarks": remarks or f"Imported candidate: {name}",
            "recommendation": "Recommended" if score >= 75 else "Consider",
            "fit_verdict": remarks or "Imported via spreadsheet",
            "source_platform": "Excel/CSV Spreadsheet Import"
        }

        if existing:
            existing.name = name or existing.name
            existing.scorecard = {**(existing.scorecard or {}), **scorecard}
            existing.status = status
            updated_count += 1
        else:
            cand_email = email if (email and "@" in email) else f"imported_{datetime.now().strftime('%M%S')}_{imported_count}@octaos.internal"
            new_cand = models.Candidate(
                tenant_id=tenant_id,
                name=name or f"Candidate {imported_count + 1}",
                email=cand_email,
                role="Imported Applicant",
                scorecard=scorecard,
                status=status
            )
            db.add(new_cand)
            imported_count += 1

    db.commit()

    return {
        "status": "success",
        "message": f"Successfully processed spreadsheet. Imported {imported_count} new candidate(s), updated {updated_count} existing candidate(s).",
        "imported_count": imported_count,
        "updated_count": updated_count,
        "total_rows": len(rows_data)
    }


from app.api.v1.endpoints.hr_lifecycle import router as hr_lifecycle_router
router.include_router(hr_lifecycle_router)


