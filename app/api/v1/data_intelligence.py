"""Governed upload-to-report workflow for business data analysis.

Uploads are processed in-memory with strict limits. Raw files are not persisted by
this endpoint; only hashes, extraction metadata, profiles and evidence-linked findings
are stored. Unsupported or unreadable files remain explicitly unavailable.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
import statistics
import zipfile
from datetime import datetime, timezone
from pathlib import PurePosixPath
from typing import Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError

from app.api.deps import get_current_user, get_db
from app.models.data_intelligence import DataAnalysisArtifact, DataAnalysisRun

router = APIRouter(prefix="/data-intelligence", tags=["orchestra-data-intelligence"])
MAX_FILES = int(os.getenv("DATA_INTELLIGENCE_MAX_FILES", "100"))
MAX_FILE_BYTES = int(os.getenv("DATA_INTELLIGENCE_MAX_FILE_BYTES", str(25 * 1024 * 1024)))
MAX_TOTAL_BYTES = int(os.getenv("DATA_INTELLIGENCE_MAX_TOTAL_BYTES", str(100 * 1024 * 1024)))
MAX_ZIP_EXPANDED_BYTES = int(os.getenv("DATA_INTELLIGENCE_MAX_ZIP_EXPANDED_BYTES", str(200 * 1024 * 1024)))
MAX_ROWS = int(os.getenv("DATA_INTELLIGENCE_MAX_ROWS", "100000"))
MAX_TEXT_CHARS = 2_000_000
SUPPORTED = {".csv", ".tsv", ".json", ".jsonl", ".ndjson", ".txt", ".md", ".log",
             ".pdf", ".docx", ".xlsx", ".xlsm", ".pptx", ".zip"}


def _safe_name(name: str) -> str:
    # Never trust client paths. ZIP entries are normalized separately.
    return PurePosixPath((name or "upload.bin").replace("\\", "/")).name[:512] or "upload.bin"


def _json_safe(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, dict):
        return {str(k): _json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(v) for v in value]
    return str(value)


def _extract_zip(data: bytes, source_name: str) -> list[tuple[str, bytes]]:
    found: list[tuple[str, bytes]] = []
    expanded = 0
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            for item in archive.infolist():
                if item.is_dir():
                    continue
                path = PurePosixPath(item.filename.replace("\\", "/"))
                if path.is_absolute() or ".." in path.parts:
                    raise HTTPException(400, detail=f"Unsafe path in archive: {item.filename[:100]}")
                if item.flag_bits & 0x1:
                    raise HTTPException(400, detail="Encrypted ZIP archives are not supported.")
                expanded += item.file_size
                if expanded > MAX_ZIP_EXPANDED_BYTES:
                    raise HTTPException(413, detail="Expanded archive exceeds configured size limit.")
                if item.file_size > MAX_FILE_BYTES:
                    raise HTTPException(413, detail=f"Archive entry exceeds per-file limit: {path.name}")
                if PurePosixPath(path.name).suffix.lower() not in SUPPORTED - {".zip"}:
                    continue
                found.append((f"{_safe_name(source_name)}/{path.as_posix()}", archive.read(item)))
                if len(found) > MAX_FILES:
                    raise HTTPException(413, detail="Archive contains too many supported files.")
    except zipfile.BadZipFile:
        raise HTTPException(400, detail=f"{source_name} has a .zip extension but is not a valid ZIP archive.")
    return found


def _read_table(data: bytes, ext: str) -> tuple[list[str], list[dict[str, Any]]]:
    text = data.decode("utf-8-sig", errors="replace")
    if ext in {".csv", ".tsv"}:
        dialect = "excel-tab" if ext == ".tsv" else "excel"
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        columns = [str(c or "").strip()[:256] for c in (reader.fieldnames or [])]
        rows = []
        for i, row in enumerate(reader):
            if i >= MAX_ROWS:
                break
            rows.append({str(k): (v[:20000] if isinstance(v, str) else v) for k, v in row.items() if k is not None})
        return columns, rows
    if ext in {".json", ".jsonl", ".ndjson"}:
        if ext in {".jsonl", ".ndjson"}:
            objects = [json.loads(line) for line in text.splitlines() if line.strip()][:MAX_ROWS]
        else:
            obj = json.loads(text)
            objects = obj if isinstance(obj, list) else [obj]
            objects = objects[:MAX_ROWS]
        if objects and all(isinstance(x, dict) for x in objects):
            cols = list(dict.fromkeys(k for row in objects for k in row.keys()))[:1000]
            return [str(c) for c in cols], [{str(k): _json_safe(v) for k, v in row.items()} for row in objects]
        return ["value"], [{"value": _json_safe(x)} for x in objects]
    return [], []


def _extract_document(data: bytes, ext: str) -> tuple[str, dict[str, Any]]:
    if ext in {".txt", ".md", ".log"}:
        return data.decode("utf-8-sig", errors="replace")[:MAX_TEXT_CHARS], {"method": "utf8_text"}
    if ext == ".pdf":
        import fitz
        doc = fitz.open(stream=data, filetype="pdf")
        pages = []
        for index, page in enumerate(doc):
            pages.append(page.get_text("text") or "")
            if sum(map(len, pages)) >= MAX_TEXT_CHARS:
                break
        return "\n".join(pages)[:MAX_TEXT_CHARS], {"method": "pymupdf", "pages": len(doc)}
    if ext == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows[:5000]:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)[:MAX_TEXT_CHARS], {"method": "python_docx", "paragraphs": len(doc.paragraphs), "tables": len(doc.tables)}
    if ext in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        output = []
        summary = []
        total_rows = 0
        for sheet in book.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                if total_rows >= MAX_ROWS:
                    break
                rows.append(" | ".join("" if v is None else str(v)[:1000] for v in row))
                total_rows += 1
            output.append(f"# Sheet: {sheet.title}\n" + "\n".join(rows))
            summary.append({"sheet": sheet.title, "rows_extracted": len(rows)})
        return "\n".join(output)[:MAX_TEXT_CHARS], {"method": "openpyxl", "sheets": summary}
    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        parts = []
        for i, slide in enumerate(prs.slides):
            text = [shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text]
            parts.append(f"# Slide {i + 1}\n" + "\n".join(text))
        return "\n".join(parts)[:MAX_TEXT_CHARS], {"method": "python_pptx", "slides": len(prs.slides)}
    raise ValueError("unsupported_format")


def _profile(columns: list[str], rows: list[dict[str, Any]]) -> dict[str, Any]:
    total = len(rows)
    fields = []
    for column in columns[:200]:
        values = [row.get(column) for row in rows]
        present = [v for v in values if v not in (None, "")]
        numeric = []
        for value in present:
            try:
                if isinstance(value, bool):
                    continue
                number = float(str(value).replace(",", "").strip())
                if math.isfinite(number):
                    numeric.append(number)
            except (ValueError, TypeError):
                pass
        field = {
            "name": column, "non_null": len(present), "null_count": total - len(present),
            "null_rate": round((total - len(present)) / total, 4) if total else None,
            "distinct_count": len({str(v) for v in present}),
            "sample_values": [_json_safe(v) for v in present[:3]],
        }
        if numeric and len(numeric) >= max(3, int(len(present) * 0.7)):
            field["type_inferred"] = "numeric"
            field["min"] = min(numeric)
            field["max"] = max(numeric)
            field["mean"] = round(statistics.fmean(numeric), 6)
            field["median"] = statistics.median(numeric)
            field["numeric_values"] = len(numeric)
        else:
            field["type_inferred"] = "text" if present else "empty"
        fields.append(field)
    return {"row_count": total, "column_count": len(columns), "columns": fields}


def _findings(name: str, profile: dict[str, Any], artifact_id: str) -> list[dict[str, Any]]:
    findings = []
    rows = profile["row_count"]
    for field in profile["columns"]:
        if rows and field.get("null_rate", 0) >= 0.2:
            findings.append({
                "id": f"F-{len(findings)+1:03d}", "severity": "medium" if field["null_rate"] < .5 else "high",
                "title": f"Missing values in {field['name']}", "observation": f"{field['null_count']} of {rows} rows ({field['null_rate']:.1%}) are empty.",
                "why_it_matters": "Missing values can bias aggregates, reduce usable records, and break downstream joins or models.",
                "recommended_action": "Confirm whether blanks mean unknown, not applicable, or a collection failure; define an explicit handling rule before filling or dropping values.",
                "evidence": [{"artifact_id": artifact_id, "file": name, "field": field["name"], "metric": "null_rate", "value": field["null_rate"]}],
                "scope": "FILE_VERIFIED"
            })
        if field.get("type_inferred") == "numeric" and field.get("min") == field.get("max") and rows >= 3:
            findings.append({
                "id": f"F-{len(findings)+1:03d}", "severity": "low", "title": f"Constant numeric field: {field['name']}",
                "observation": f"The observed numeric values for {field['name']} are all {field['min']}.",
                "why_it_matters": "A constant field may be intentional, incorrectly populated, or provide no analytical separation.",
                "recommended_action": "Validate the source definition and ingestion mapping before using this field in segmentation or modeling.",
                "evidence": [{"artifact_id": artifact_id, "file": name, "field": field["name"], "metric": "min_equals_max", "value": field["min"]}],
                "scope": "FILE_VERIFIED"
            })
    if rows == 0:
        findings.append({"id":"F-001","severity":"high","title":"No tabular records extracted","observation":"The file was read but produced no tabular records.","why_it_matters":"No row-level conclusions can be drawn from an empty or unsupported structure.","recommended_action":"Verify file contents and export format, then upload a valid table or supported document.","evidence":[{"artifact_id":artifact_id,"file":name,"metric":"row_count","value":0}],"scope":"FILE_VERIFIED"})
    return findings


@router.post("/runs", status_code=status.HTTP_201_CREATED)
async def create_analysis_run(
    files: list[UploadFile] = File(..., description="Upload multiple files; upload a ZIP to represent a folder."),
    title: str = Form("Business data analysis"),
    idempotency_key: str | None = Form(None),
    current_user=Depends(get_current_user),
    db: Session = Depends(get_db),
):
    tenant_id = getattr(current_user, "tenant_id", None)
    user_id = getattr(current_user, "id", None)
    if not tenant_id or not user_id:
        raise HTTPException(403, detail="An authenticated user with an active tenant is required.")
    if not files:
        raise HTTPException(400, detail="Upload at least one file.")
    if len(files) > MAX_FILES:
        raise HTTPException(413, detail=f"Maximum {MAX_FILES} uploaded files per run.")
    if idempotency_key and (len(idempotency_key) > 128 or any(ord(c) < 32 for c in idempotency_key)):
        raise HTTPException(400, detail="Invalid idempotency key.")

    incoming: list[tuple[str, bytes]] = []
    total_bytes = 0
    for upload in files:
        name = _safe_name(upload.filename or "upload.bin")
        data = await upload.read(MAX_FILE_BYTES + 1)
        if len(data) > MAX_FILE_BYTES:
            raise HTTPException(413, detail=f"{name} exceeds the {MAX_FILE_BYTES}-byte per-file limit.")
        total_bytes += len(data)
        if total_bytes > MAX_TOTAL_BYTES:
            raise HTTPException(413, detail=f"Upload exceeds total limit of {MAX_TOTAL_BYTES} bytes.")
        ext = PurePosixPath(name).suffix.lower()
        if ext == ".zip":
            incoming.extend(_extract_zip(data, name))
        else:
            incoming.append((name, data))
    if len(incoming) > MAX_FILES:
        raise HTTPException(413, detail=f"Expanded upload exceeds {MAX_FILES} files.")
    if not incoming:
        raise HTTPException(400, detail="No supported files found. Supported: CSV, TSV, JSON, JSONL, TXT, MD, LOG, PDF, DOCX, XLSX, XLSM, PPTX, ZIP.")
    if sum(len(data) for _, data in incoming) > MAX_TOTAL_BYTES:
        raise HTTPException(413, detail="Expanded files exceed the total byte limit.")

    hashes = [{"name": name, "sha256": hashlib.sha256(data).hexdigest(), "size": len(data)} for name, data in incoming]
    fingerprint = hashlib.sha256(json.dumps(hashes, sort_keys=True, separators=(",", ":")).encode()).hexdigest()
    if idempotency_key:
        prior = db.query(DataAnalysisRun).filter_by(tenant_id=str(tenant_id), idempotency_key=idempotency_key).first()
        if prior:
            if prior.request_fingerprint != fingerprint:
                raise HTTPException(409, detail="Idempotency key was already used with a different upload payload.")
            return {"run": _serialize_run(prior), "idempotent_replay": True}

    run = DataAnalysisRun(
        tenant_id=str(tenant_id), created_by_user_id=str(user_id), title=title.strip()[:255] or "Business data analysis",
        status="processing", idempotency_key=idempotency_key, request_fingerprint=fingerprint,
        file_count=len(incoming), total_bytes=sum(len(data) for _, data in incoming),
        started_at=datetime.now(timezone.utc), profile={}, findings=[], report={},
    )
    db.add(run)
    db.flush()
    all_findings = []
    artifact_summaries = []
    combined_profiles = []
    try:
        for name, data in incoming:
            ext = PurePosixPath(name).suffix.lower()
            sha = hashlib.sha256(data).hexdigest()
            artifact_id = __import__("uuid").uuid4().hex
            artifact = DataAnalysisArtifact(
                id=artifact_id, run_id=run.id, tenant_id=str(tenant_id), filename=name[:512],
                media_type=next((x for x in files if _safe_name(x.filename or "") == name), None).content_type if any(_safe_name(x.filename or "") == name for x in files) else "application/octet-stream",
                sha256=sha, byte_size=len(data), extraction_status="unavailable", extraction_summary={},
            )
            try:
                if ext in {".csv", ".tsv", ".json", ".jsonl", ".ndjson"}:
                    columns, rows = _read_table(data, ext)
                    profile = _profile(columns, rows)
                    artifact.extraction_status = "completed"
                    artifact.extraction_summary = {"kind": "table", "row_count": len(rows), "columns": columns[:200]}
                    found = _findings(name, profile, artifact_id)
                    combined_profiles.append({"artifact_id": artifact_id, "file": name, "kind": "table", **profile})
                elif ext in {".txt", ".md", ".log", ".pdf", ".docx", ".xlsx", ".xlsm", ".pptx"}:
                    content, meta = _extract_document(data, ext)
                    artifact.extraction_status = "completed" if content.strip() else "empty"
                    artifact.extraction_summary = {"kind": "document", "character_count": len(content), **meta}
                    found = []
                    combined_profiles.append({"artifact_id": artifact_id, "file": name, "kind": "document", "character_count": len(content), "extraction": meta,
                                              "preview": content[:1200]})
                    if not content.strip():
                        found.append({"id":f"F-{len(all_findings)+1:03d}","severity":"medium","title":f"No text extracted from {name}","observation":"The parser returned no text. Scanned pages or embedded images may require OCR and are not claimed as analyzed.","why_it_matters":"Text-only extraction can miss charts, screenshots, scanned PDFs and visual-only content.","recommended_action":"Run OCR or export the source to a text-accessible format before relying on content conclusions.","evidence":[{"artifact_id":artifact_id,"file":name,"metric":"extracted_characters","value":0}],"scope":"FILE_VERIFIED"})
                    found = found if content.strip() else found
                else:
                    found = [{"id":f"F-{len(all_findings)+1:03d}","severity":"high","title":f"Unsupported file type: {name}","observation":"This file format is not supported by the configured extractor.","why_it_matters":"No claims can be grounded in a file that could not be parsed.","recommended_action":"Convert to a supported format or add a reviewed parser for this type.","evidence":[{"artifact_id":artifact_id,"file":name,"metric":"extraction_status","value":"unavailable"}],"scope":"FILE_VERIFIED"}]
                all_findings.extend(found)
            except Exception as exc:
                artifact.extraction_status = "failed"
                artifact.extraction_summary = {"error_code": type(exc).__name__, "message": "Extraction failed; contents were not analyzed."}
                all_findings.append({"id":f"F-{len(all_findings)+1:03d}","severity":"high","title":f"Could not analyze {name}","observation":f"Extraction failed ({type(exc).__name__}); no content-level conclusions were generated.","why_it_matters":"An unreadable source can leave material gaps in the report.","recommended_action":"Validate the file and retry; review parser logs without exposing sensitive content.","evidence":[{"artifact_id":artifact_id,"file":name,"metric":"extraction_status","value":"failed"}],"scope":"FILE_VERIFIED"})
            db.add(artifact)
            artifact_summaries.append({"id":artifact_id,"filename":name,"sha256":sha,"size_bytes":len(data),"status":artifact.extraction_status,"summary":artifact.extraction_summary})
        run.profile = {"artifacts": combined_profiles, "files_received": len(incoming), "files_parsed": sum(1 for a in artifact_summaries if a["status"] == "completed")}
        run.findings = all_findings
        run.status = "completed" if any(a["status"] == "completed" for a in artifact_summaries) else "partial"
        run.completed_at = datetime.now(timezone.utc)
        run.report = {
            "schema_version": "1.0", "run_id": run.id, "title": run.title, "status": run.status,
            "generated_at": run.completed_at.isoformat(), "methodology": "Deterministic file extraction and descriptive profiling. No external data, LLM inference, causality, or business KPI values are invented.",
            "coverage": {"files_received": len(incoming), "files_parsed": run.profile["files_parsed"], "files_unavailable_or_failed": len(incoming)-run.profile["files_parsed"], "bytes_processed": run.total_bytes},
            "artifacts": artifact_summaries, "profiles": combined_profiles, "findings": all_findings,
            "limitations": ["Scanned pages and visual-only charts may not be extracted.", "Statistical associations are not causal explanations.", "Recommendations are validation guidance, not proof of business impact.", "No source files are retained by this API after request processing."],
        }
        db.commit()
        db.refresh(run)
    except IntegrityError:
        db.rollback()
        # A concurrent request may have won the unique tenant/idempotency-key race.
        if idempotency_key:
            prior = db.query(DataAnalysisRun).filter_by(tenant_id=str(tenant_id), idempotency_key=idempotency_key).first()
            if prior:
                if prior.request_fingerprint != fingerprint:
                    raise HTTPException(409, detail="Idempotency key was already used with a different upload payload.")
                return {"run": _serialize_run(prior), "idempotent_replay": True}
        raise HTTPException(409, detail="A conflicting analysis write was detected; retry with a new idempotency key.")
    except Exception:
        db.rollback()
        raise
    return {"run": _serialize_run(run), "idempotent_replay": False}


def _serialize_run(run: DataAnalysisRun) -> dict[str, Any]:
    return {"id":run.id,"title":run.title,"status":run.status,"file_count":run.file_count,"total_bytes":run.total_bytes,
            "created_at":run.created_at.isoformat() if run.created_at else None,
            "completed_at":run.completed_at.isoformat() if run.completed_at else None,"finding_count":len(run.findings or [])}


def _owned_run(db: Session, run_id: str, tenant_id: str) -> DataAnalysisRun:
    run = db.query(DataAnalysisRun).filter(DataAnalysisRun.id == run_id, DataAnalysisRun.tenant_id == str(tenant_id)).first()
    if not run:
        raise HTTPException(404, detail="Analysis run not found.")
    return run


@router.get("/runs")
def list_analysis_runs(limit: int = 20, offset: int = 0, current_user=Depends(get_current_user), db: Session = Depends(get_db)):
    tenant_id = getattr(current_user, "tenant_id", None)
    if not tenant_id:
        raise HTTPException(403, detail="An active tenant is required.")
    limit = max(1, min(limit, 100))
    offset = max(0, offset)
    rows = db.query(DataAnalysisRun).filter_by(tenant_id=str(tenant_id)).order_by(DataAnalysisRun.created_at.desc()).offset(offset).limit(limit).all()
    return {"items":[_serialize_run(row) for row in rows],"limit":limit,"offset":offset}


@router.get("/runs/{run_id}")
def get_analysis_run(run_id: str, current_user=Depends(get_current_user), db: Session = Depends(get_db)):
    tenant_id = getattr(current_user, "tenant_id", None)
    if not tenant_id:
        raise HTTPException(403, detail="An active tenant is required.")
    run = _owned_run(db, run_id, str(tenant_id))
    return {"run":_serialize_run(run),"profile":run.profile,"findings":run.findings,"report":run.report}


@router.get("/runs/{run_id}/report.json")
def export_analysis_json(run_id: str, current_user=Depends(get_current_user), db: Session = Depends(get_db)):
    tenant_id = getattr(current_user, "tenant_id", None)
    if not tenant_id:
        raise HTTPException(403, detail="An active tenant is required.")
    run = _owned_run(db, run_id, str(tenant_id))
    if not run.report:
        raise HTTPException(409, detail="Report is not ready.")
    return run.report


@router.get("/runs/{run_id}/report.md")
def export_analysis_markdown(run_id: str, current_user=Depends(get_current_user), db: Session = Depends(get_db)):
    from fastapi.responses import PlainTextResponse
    tenant_id = getattr(current_user, "tenant_id", None)
    if not tenant_id:
        raise HTTPException(403, detail="An active tenant is required.")
    run = _owned_run(db, run_id, str(tenant_id))
    if not run.report:
        raise HTTPException(409, detail="Report is not ready.")
    report = run.report
    lines = [f"# {report.get('title','Data Intelligence Report')}", "", f"- Run ID: `{run.id}`", f"- Status: {run.status}",
             f"- Generated: {report.get('generated_at')}", "", "## Coverage", ""]
    for key, value in report.get("coverage", {}).items():
        lines.append(f"- **{key.replace('_',' ').title()}:** {value}")
    lines += ["", "## File profiles", ""]
    for item in report.get("profiles", []):
        lines += [f"### {item.get('file')}", "", f"- Kind: {item.get('kind')}", f"- Rows: {item.get('row_count', 'N/A')}", f"- Columns: {item.get('column_count', 'N/A')}"]
        for field in item.get("columns", []):
            lines.append(f"- **{field.get('name')}** — {field.get('type_inferred')}; null rate: {field.get('null_rate', 'N/A')}; distinct: {field.get('distinct_count')}")
        if item.get("preview"):
            lines += ["", item["preview"][:1200]]
        lines.append("")
    lines += ["## Findings", ""]
    for finding in report.get("findings", []):
        lines += [f"### [{finding.get('severity','info').upper()}] {finding.get('title')}", "", f"**Observed:** {finding.get('observation')}", "", f"**Why it matters:** {finding.get('why_it_matters')}", "", f"**Recommended action:** {finding.get('recommended_action')}", "", f"**Evidence:** {json.dumps(finding.get('evidence', []), ensure_ascii=False)}", ""]
    lines += ["## Limitations", ""] + [f"- {item}" for item in report.get("limitations", [])]
    return PlainTextResponse("\n".join(lines), media_type="text/markdown; charset=utf-8",
                             headers={"Content-Disposition": f'attachment; filename="data-intelligence-{run.id}.md"'})
